import os
import sys
import unittest
from pathlib import Path

os.environ.setdefault('MONGO_URL', 'mongodb://localhost:27017')
os.environ.setdefault('SKIP_STARTUP_INIT', '1')
sys.path.insert(0, str(Path(__file__).resolve().parent))

import asyncio
from io import BytesIO

from docx import Document
from fastapi import HTTPException
from starlette.datastructures import UploadFile

import main  # noqa: E402


class ConceptoImportParsingTests(unittest.TestCase):
    def test_detects_headers_regardless_of_order_and_naming(self):
        rows = [
            ['Precio Unitario', 'Descripción', 'Cantidad', 'Unidad'],
            [60, 'Tubería PVC 4"', 100, 'ml'],
            [200, 'Conexiones', 20, 'pza'],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows)

        self.assertEqual(len(items), 2)
        self.assertEqual(items[0], {'description': 'Tubería PVC 4"', 'unit': 'ml', 'quantity': 100.0, 'unitPrice': 60.0})
        self.assertEqual(items[1], {'description': 'Conexiones', 'unit': 'pza', 'quantity': 20.0, 'unitPrice': 200.0})
        self.assertEqual(warnings, [])

    def test_falls_back_to_fixed_order_when_no_recognizable_headers(self):
        rows = [
            ['Excavación a mano', 'm3', 50, 120],
            ['Relleno compactado', 'm3', 30, 90],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows)

        self.assertEqual(len(items), 2)
        self.assertEqual(items[0]['description'], 'Excavación a mano')
        self.assertEqual(items[0]['quantity'], 50.0)
        self.assertEqual(items[0]['unitPrice'], 120.0)
        self.assertTrue(any('orden' in w.lower() for w in warnings))

    def test_parses_currency_formatted_numbers(self):
        rows = [
            ['Concepto', 'Unidad', 'Cantidad', 'Precio Unitario'],
            ['Acabados', 'lote', '1,500.00', '$3,200.50'],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows)

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]['quantity'], 1500.0)
        self.assertEqual(items[0]['unitPrice'], 3200.5)

    def test_skips_total_row_without_numeric_quantity_or_price(self):
        rows = [
            ['Concepto', 'Unidad', 'Cantidad', 'Precio Unitario'],
            ['Plomería', 'lote', 10, 100],
            ['TOTAL', '', '', ''],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows)

        self.assertEqual(len(items), 1)
        # un renglón TOTAL se reconoce como tal: no cuenta como renglón inválido
        self.assertFalse(any('omitieron' in w.lower() for w in warnings))

    def test_skips_rows_missing_description(self):
        rows = [
            ['Concepto', 'Unidad', 'Cantidad', 'Precio Unitario'],
            ['', 'lote', 10, 100],
            ['Herrería', 'lote', 5, 50],
        ]
        items, _warnings = main.parse_concepto_rows_from_table(rows)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]['description'], 'Herrería')

    def test_empty_table_returns_no_items_no_warnings(self):
        items, warnings = main.parse_concepto_rows_from_table([])
        self.assertEqual(items, [])
        self.assertEqual(warnings, [])

    def test_low_confidence_flag_adds_warning_when_items_found(self):
        rows = [
            ['Concepto', 'Unidad', 'Cantidad', 'Precio Unitario'],
            ['Impermeabilizante', 'm2', 80, 45],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows, low_confidence=True)
        self.assertEqual(len(items), 1)
        self.assertTrue(any('baja confianza' in w.lower() for w in warnings))

    def test_pdf_text_fallback_splits_on_multiple_spaces(self):
        # Mirrors the heuristic used for scanned-looking PDFs with no
        # extractable table grid: split each text line on runs of 2+ spaces.
        import re
        text = "Concepto        Unidad   Cantidad   Precio Unitario\nMuro de block   m2       120        180.00\n"
        text_rows = [re.split(r"\s{2,}", line.strip()) for line in text.splitlines() if line.strip()]
        items, warnings = main.parse_concepto_rows_from_table(text_rows, low_confidence=True)

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]['description'], 'Muro de block')
        self.assertEqual(items[0]['quantity'], 120.0)
        self.assertEqual(items[0]['unitPrice'], 180.0)
        self.assertTrue(any('baja confianza' in w.lower() for w in warnings))


class ConceptoHeaderDetectionTests(unittest.TestCase):
    def test_normalize_concepto_header_strips_accents_and_symbols(self):
        self.assertEqual(main.normalize_concepto_header('Descripción'), 'descripcion')
        self.assertEqual(main.normalize_concepto_header('Precio Unitario'), 'preciounitario')
        self.assertEqual(main.normalize_concepto_header(' Cantidad '), 'cantidad')

    def test_detect_concepto_header_row_requires_all_core_columns(self):
        rows = [
            ['Notas'],
            ['Concepto', 'Cantidad', 'Precio Unitario'],
            ['Pintura', 10, 50],
        ]
        row_idx, header_index = main.detect_concepto_header_row(rows)
        self.assertEqual(row_idx, 1)
        self.assertEqual(header_index, {'concepto': 0, 'cantidad': 1, 'preciounitario': 2})

    def test_detect_concepto_header_row_returns_none_when_incomplete(self):
        rows = [['Concepto', 'Unidad'], ['Pintura', 'lote']]
        row_idx, header_index = main.detect_concepto_header_row(rows)
        self.assertIsNone(row_idx)
        self.assertEqual(header_index, {})


def _docx_bytes(build):
    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    return table


class DocxImportTests(unittest.TestCase):
    def test_reads_concepts_from_docx_table_with_headers(self):
        data = _docx_bytes(lambda d: (
            d.add_paragraph('Presupuesto de herrería'),
            _add_table(d, [
                ['Concepto', 'Unidad', 'Cantidad', 'P.U.', 'Importe'],
                ['Reja perimetral', 'ml', '40', '$1,250.50', '$50,020.00'],
                ['Puerta', 'pza', '2', '8,000', '16,000'],
                ['TOTAL', '', '', '', '66,020.00'],
            ]),
        ))
        items, warnings = main.extract_concepto_rows_from_docx_bytes(data)
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0], {'description': 'Reja perimetral', 'unit': 'ml', 'quantity': 40.0, 'unitPrice': 1250.5})
        self.assertEqual(items[1]['unitPrice'], 8000.0)
        self.assertFalse(any('omitieron' in w for w in warnings))  # fila TOTAL reconocida, no es un error

    def test_ignores_tables_without_concept_columns(self):
        data = _docx_bytes(lambda d: (
            _add_table(d, [['Cliente', 'Obra'], ['Juan', 'Casa 1']]),
            _add_table(d, [['Descripción', 'Unidad', 'Cantidad', 'Precio unitario'], ['Muro', 'm2', '10', '300']]),
        ))
        items, warnings = main.extract_concepto_rows_from_docx_bytes(data)
        self.assertEqual([i['description'] for i in items], ['Muro'])
        self.assertTrue(any('ignoraron 1 tabla' in w for w in warnings))

    def test_merged_cells_do_not_shift_columns(self):
        def build(d):
            table = _add_table(d, [
                ['Concepto', 'Unidad', 'Cantidad', 'Precio unitario'],
                ['Aplanado', 'm2', '20', '150'],
                ['Nota', 'x', 'x', 'x'],
            ])
            table.cell(2, 0).merge(table.cell(2, 3))
        items, _ = main.extract_concepto_rows_from_docx_bytes(_docx_bytes(build))
        self.assertEqual([i['description'] for i in items], ['Aplanado'])

    def test_docx_without_tables_returns_nothing(self):
        items, _ = main.extract_concepto_rows_from_docx_bytes(_docx_bytes(lambda d: d.add_paragraph('Solo texto')))
        self.assertEqual(items, [])

    def _call_endpoint(self, filename, content):
        upload = UploadFile(file=BytesIO(content), filename=filename)
        return asyncio.run(main.import_estimation_conceptos(file=upload, user={'role': 'ADMIN'}))

    def test_endpoint_accepts_docx(self):
        data = _docx_bytes(lambda d: _add_table(d, [
            ['Concepto', 'Unidad', 'Cantidad', 'Precio unitario'], ['Muro', 'm2', '10', '300'],
        ]))
        result = self._call_endpoint('Presupuesto.DOCX', data)
        self.assertEqual(result['sourceType'], 'docx')
        self.assertEqual(result['items'][0]['quantity'], 10.0)

    def test_endpoint_rejects_legacy_doc_with_helpful_message(self):
        with self.assertRaises(HTTPException) as ctx:
            self._call_endpoint('viejo.doc', b'\xd0\xcf\x11\xe0')
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn('.docx', ctx.exception.detail)

    def test_endpoint_rejects_corrupt_docx(self):
        with self.assertRaises(HTTPException) as ctx:
            self._call_endpoint('roto.docx', b'not a zip')
        self.assertEqual(ctx.exception.status_code, 400)

    def test_endpoint_reports_docx_without_concepts(self):
        data = _docx_bytes(lambda d: d.add_paragraph('Sin tabla'))
        with self.assertRaises(HTTPException) as ctx:
            self._call_endpoint('vacio.docx', data)
        self.assertEqual(ctx.exception.status_code, 422)


class ContractorBudgetFormatTests(unittest.TestCase):
    """Formato real de presupuestos de contratistas: montos con centavos como
    superíndice (4935ºº), encabezado propio (MUEBLES), columnas en otro orden,
    secciones sin encabezado y una tabla final de resumen."""

    def _doc(self):
        def build(d):
            d.add_paragraph('PRESUPUESTO DE INSTALACIONES')
            d.add_paragraph('INSTALACION DE TUBERIA PARA MUEBLES')
            _add_table(d, [
                ['MUEBLES', 'CANTIDAD', 'UNIDAD', 'P/UNITARIO', 'IMPORTE'],
                ['W.C.', '24', 'PZZAS', '$ 4935ºº', '$ 118440ºº'],
                ['REGADERAS', '18', 'PZZAS', '$ 4200ºº', '$   75600ºº'],
                ['', '', '', 'TOTAL', '$ 194040ºº'],
            ])
            d.add_paragraph('INSTALACION DE BAJADAS')
            _add_table(d, [
                ['B.A.N.', '5', 'PZZAS', '$ 7700ºº', '$ 38521ºº'],
                ['LINEA GENERAL DE DRENAJE', '1', 'LOTE', '$ 60911ºº', '$ 60911º'],
            ])
            d.add_paragraph('COLOCACION DE MUEBLES')
            _add_table(d, [
                ['W.C.', '24', 'PZZAS', '$ 450ºº', '$ 10800ºº'],
                ['', '', '', 'TOTAL', '$ 10800ºº'],
            ])
            d.add_paragraph('GENERALES')
            _add_table(d, [
                ['INSTALACION DE TUBERIA PARA MUEBLES', '$ 194040ºº'],
                ['TOTAL', '$ 999999ºº'],
            ])
        return _docx_bytes(build)

    def test_parses_superscript_cents_and_custom_header(self):
        for raw, expected in (('$ 4935ºº', 4935.0), ('$ 19650º', 19650.0), ('$   75600ºº', 75600.0), ('$ 1,250.50', 1250.5)):
            self.assertEqual(main.parse_concepto_number(raw), expected)

    def test_unrecognized_first_column_header_is_used_as_concepto(self):
        rows = [
            ['MUEBLES', 'CANTIDAD', 'UNIDAD', 'P/UNITARIO', 'IMPORTE'],
            ['W.C.', '24', 'PZZAS', '$ 4935ºº', '$ 118440ºº'],
        ]
        items, _ = main.parse_concepto_rows_from_table(rows)
        self.assertEqual(items, [{'description': 'W.C.', 'unit': 'PZZAS', 'quantity': 24.0, 'unitPrice': 4935.0}])

    def test_reads_every_section_even_without_repeating_the_header(self):
        items, warnings = main.extract_concepto_rows_from_docx_bytes(self._doc())
        descriptions = [item['description'] for item in items]
        self.assertEqual(len(items), 5)
        self.assertIn('REGADERAS', descriptions)
        self.assertIn('LINEA GENERAL DE DRENAJE', descriptions)
        by_name = {item['description']: item for item in items}
        self.assertEqual(by_name['B.A.N.']['quantity'], 5.0)
        self.assertEqual(by_name['B.A.N.']['unitPrice'], 7700.0)
        self.assertEqual(by_name['B.A.N.']['unit'], 'PZZAS')

    def test_each_titled_table_becomes_a_group(self):
        items, _ = main.extract_concepto_rows_from_docx_bytes(self._doc())
        groups = {item['description'] + '|' + item.get('group', ''): item for item in items}
        self.assertIn('W.C.|INSTALACION DE TUBERIA PARA MUEBLES', groups)
        self.assertIn('W.C.|COLOCACION DE MUEBLES', groups)
        self.assertEqual(groups['B.A.N.|INSTALACION DE BAJADAS']['group'], 'INSTALACION DE BAJADAS')
        # el nombre del concepto ya no se modifica
        self.assertTrue(all('(' not in item['description'] for item in items))

    def test_title_only_rows_in_a_spreadsheet_open_groups(self):
        rows = [
            ['Concepto', 'Unidad', 'Cantidad', 'Precio Unitario'],
            ['BAJADAS', '', '', ''],
            ['B.A.N.', 'pza', 5, 7700],
            ['CUARTO DE BOMBAS', '', '', ''],
            ['Conexión de bomba', 'pza', 1, 4400],
            ['TOTAL', '', '', ''],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows)
        self.assertEqual([i.get('group') for i in items], ['BAJADAS', 'CUARTO DE BOMBAS'])
        self.assertFalse(any('omitieron' in w for w in warnings))

    def test_group_column_is_read_when_present(self):
        rows = [
            ['Grupo', 'Concepto', 'Unidad', 'Cantidad', 'P.U.'],
            ['Drenaje', 'Registro', 'pza', 2, 900],
            ['Hidráulica', 'Toma', 'pza', 3, 300],
        ]
        items, _ = main.parse_concepto_rows_from_table(rows)
        self.assertEqual([i['group'] for i in items], ['Drenaje', 'Hidráulica'])

    def test_summary_table_is_ignored_but_its_total_is_checked(self):
        _, warnings = main.extract_concepto_rows_from_docx_bytes(self._doc())
        self.assertTrue(any('Se ignoraron 1 tabla' in w for w in warnings))
        self.assertTrue(any('$999,999.00' in w and 'no coincide' in w for w in warnings))

    def test_warns_when_stated_amount_differs_from_quantity_times_price(self):
        _, warnings = main.extract_concepto_rows_from_docx_bytes(self._doc())
        joined = ' '.join(warnings)
        self.assertIn('B.A.N.', joined)  # 5 x 7700 = 38,500 pero el documento dice 38,521
        self.assertIn('$38,521.00', joined)

    def test_matching_section_total_does_not_warn(self):
        rows = [
            ['MUEBLES', 'CANTIDAD', 'UNIDAD', 'P/UNITARIO', 'IMPORTE'],
            ['W.C.', '24', 'PZZAS', '$ 4935ºº', '$ 118440ºº'],
            ['', '', '', 'TOTAL', '$ 118440ºº'],
        ]
        _, warnings = main.parse_concepto_rows_from_table(rows)
        self.assertEqual(warnings, [])


if __name__ == '__main__':
    unittest.main()


class ContractorLetterheadTests(unittest.TestCase):
    """Formatos de contratista con membrete: la tabla empieza muchos renglones abajo."""

    def test_header_row_far_below_the_letterhead_is_detected(self):
        rows = [
            ['Héctor Morales Medina.', None, None, None, None, None, None, None],
            ['Contratista', None, None, None, None, None, None, '29/09/2026'],
            ['Obra:', None, 'Calderón de la Barca.', None, None, None, None, None],
            ['Contrato:', None, 'Yesería.', None, None, None, None, None],
            [None, 'Arq. Rafael Marcos.', None, None, None, None, None, None],
            [None, None, None, 'DEPARTAMENTO 201.', None, None, None, None],
            [None, None, 'Por este conducto presente el siguiente presupuesto a su consideración.', None, None, None, None, None],
            [None] * 8,
            [None] * 8,
            ['No.', 'Concepto', None, None, 'Unidad', 'Cantidad', 'Precio', 'Total'],
            [1.0, 'Aplanados de yeso en muros.', None, None, 'm2.', 480.6, 115.0, ' $ 55,269.00 '],
            [2.0, 'Emboquillados en muros.', None, None, 'm.', 109.0, 57.5, ' $ 6,267.50 '],
            [3.0, 'Refuerzo con metal desplegado cal. 500', None, None, 'm.', 82.0, 90.0, ' $ 7,380.00 '],
            [None] * 8,
            [None, None, None, None, None, None, 'TOTAL', ' $ 68,916.50 '],
            [None, 'Notas.', None, None, None, None, None, None],
            [None, None, 'No incluye I.V.A.', None, None, None, None, None],
        ]
        items, warnings = main.parse_concepto_rows_from_table(rows)
        self.assertEqual(
            [(i['description'], i['unit'], i['quantity'], i['unitPrice']) for i in items],
            [
                ('Aplanados de yeso en muros.', 'm2', 480.6, 115.0),
                ('Emboquillados en muros.', 'm', 109.0, 57.5),
                ('Refuerzo con metal desplegado cal. 500', 'm', 82.0, 90.0),
            ],
        )
        self.assertFalse(any('Se asumió' in w or 'no coincide' in w for w in warnings))
